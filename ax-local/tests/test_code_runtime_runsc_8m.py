import json
import os
import subprocess
import time
import unittest

import test_code_runtime_runsc as base
from test_code_runtime_8m import DRIVER, four_files
from test_code_runtime_rpc import fixture, runner


@unittest.skipUnless(base.NODE and base.BUNDLE and base.RUNSC,'explicit owned runsc bundle required')
class RunscEightMiBTests(unittest.TestCase):
    dirty=False
    verify_idle=classmethod(base.RunscCodeRuntimeTests.verify_idle.__func__)
    run_code=base.RunscCodeRuntimeTests.run_code
    @classmethod
    def setUpClass(cls):
        base.MOUNTS={"input":("/sandbox/input",8404992,16,0,"755"),"tmp":("/sandbox/tmp",8388608,128,65532,"700"),"output":("/sandbox/output",8400896,32,65532,"700"),"state":("/var/lib/ax-code",262144,32,0,"700")}
        base.RunscCodeRuntimeTests.setUpClass.__func__(cls)
    @classmethod
    def node(cls,*args,input=None,check=True,timeout=30):
        if args and args[0]==base.RUNSC and 'run' in args:timeout=180
        return base.RunscCodeRuntimeTests.node.__func__(cls,*args,input=input,check=check,timeout=timeout)
    def invoke(self,content,source,stage=None):
        driver=DRIVER.replace("request=json.load(sys.stdin);stage=request['stage'];raw=base64.b64decode(request['content'])", "request=json.load(sys.stdin);stage=json.loads(request['source']);raw=base64.b64decode(request['inputs']['payload'])")
        driver=driver.replace("Path.read_text=mountinfo", "assert '- 9p ' in Path('/proc/self/mountinfo').read_text()")
        started=time.monotonic()
        value=self.run_code(json.dumps(stage or fixture(content,source)),inputs={'payload':content},driver=driver)
        print('runsc_8m_elapsed_seconds='+str(round(time.monotonic()-started,3)),flush=True)
        return value
    def test_port80_runner_with_five_capabilities(self):
        driver=r'''import subprocess,time,json,urllib.request
p=subprocess.Popen(['/usr/local/bin/ax-task-runner','--task-file','/opt/ax-code/server-task.json'],stdout=subprocess.PIPE,stderr=subprocess.PIPE)
ready=False
try:
 for i in range(150):
  if p.poll() is not None:break
  try:
   with urllib.request.urlopen('http://127.0.0.1:80/readyz',timeout=.1) as r:ready=r.status==200
  except Exception:pass
  if ready:break
  time.sleep(.02)
finally:
 if p.poll() is None:p.terminate()
 out,err=p.communicate(timeout=3)
print(json.dumps({'ready':ready,'exit':p.returncode,'error':err.decode(errors='replace')[-1024:],'log':out.decode(errors='replace')[-1024:]}))
'''
        base.MOUNTS['workspace']=('/workspace',65536,16,0,'755')
        try:
            value=self.run_code('pass',driver=driver)
        finally:
            del base.MOUNTS['workspace']
        self.assertTrue(value['ready'],value)
    def test_exact_8m_rpc_and_isolation(self):
        source="""import os,socket
assert os.getuid()==65532
try:socket.socket()
except PermissionError:pass
else:raise AssertionError('network')
data=open('/input/input_1','rb').read();open('/output/result.csv','wb').write(data);print(len(data))
"""
        content=b'a'*runner.LIMIT;value=self.invoke(content,source)
        self.assertEqual(value['result']['status'],'succeeded',value)
        self.assertEqual(value['collected'],[{'size_bytes':len(content),'sha256':runner.digest(content)}])
    def test_four_unaligned_files_exact_logical_limit(self):
        content,stage=four_files();value=self.invoke(content,stage['workbench']['descriptor']['code']['source'],stage)
        self.assertEqual(value['result']['status'],'succeeded',value)
        self.assertEqual(sum(f['size_bytes'] for f in value['collected']),runner.LIMIT)
    def test_csv_and_xlsx_under_existing_limits(self):
        source="""import csv
from openpyxl import Workbook,load_workbook
from openpyxl.xml import DEFUSEDXML
assert DEFUSEDXML
amount=sum(int(r['amount']) for r in csv.DictReader(open('/input/input_1')))
w=Workbook();w.active.append(['sum',amount]);w.save('/tmp/sheet.xlsx')
assert load_workbook('/tmp/sheet.xlsx',read_only=True).active['B1'].value==amount
open('/output/result.csv','w').write('sum\\n'+str(amount)+'\\n');print(amount)
"""
        value=self.invoke(b'amount\n2\n3\n',source)
        self.assertEqual(value['result']['status'],'succeeded',value)
        self.assertEqual(value['result']['summary'],'5\n')

if __name__=='__main__':unittest.main()

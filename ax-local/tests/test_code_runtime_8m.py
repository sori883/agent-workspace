import base64
import json
import os
import subprocess
import unittest
from test_code_runtime_rpc import fixture, runner

IMAGE=os.environ.get('AX_CODE_IMAGE')
DRIVER=r'''import sys,os,json,time,base64,hashlib,subprocess
from pathlib import Path
sys.path.insert(0,'/opt/ax-code')
import runner
request=json.load(sys.stdin);stage=request['stage'];raw=base64.b64decode(request['content'])
original_read_text=Path.read_text
def mountinfo(path,*args,**kwargs):
 value=original_read_text(path,*args,**kwargs)
 if str(path)=='/proc/self/mountinfo':value=value.replace(' - tmpfs ',' - 9p ')
 return value
Path.read_text=mountinfo
pid=os.fork()
if pid==0:
 try:runner.wait()
 except BaseException:os._exit(1)
 os._exit(0)
deadline=time.monotonic()+5
while not (runner.STATE/'ready.json').exists():
 assert time.monotonic()<deadline
 time.sleep(.01)
key={k:stage['request'][k] for k in ['run_id','descriptor_sha256']}
def rpc(op,value,ok=True):
 arg=base64.b64encode(runner.canonical(value)).decode()
 p=subprocess.run(['python3','-I','-B','/opt/ax-code/runner.py',op,arg],capture_output=True,timeout=10)
 assert (p.returncode==0)==ok,(op,p.stdout[:200],p.stderr[:200])
 return json.loads(p.stdout)
rpc('stage-begin',stage);rpc('stage-begin',stage)
offset=0
for f in stage['workbench']['descriptor']['inputs']:
 file_content=raw[offset:offset+f['size_bytes']];offset+=f['size_bytes']
 for i in range((len(file_content)+32767)//32768):
  content=file_content[i*32768:(i+1)*32768]
  v={**key,'alias':f['alias'],'index':i,'content_base64':base64.b64encode(content).decode(),'sha256':hashlib.sha256(content).hexdigest()}
  rpc('stage-chunk',v)
rpc('stage-seal',key);rpc('stage-seal',key)
rpc('code-start',key);rpc('code-start',key,False)
_,exit_status=os.waitpid(pid,0);assert exit_status==0,exit_status
state=rpc('code-status',key);result=rpc('code-collect',key)
assert state['state']=='finished' and state['attempted'] and state['result']==result['result']
collected=[]
if result['result']['status']=='succeeded':
 manifest=rpc('output-manifest',key)
 expected=manifest.pop('manifest_sha256');assert runner.digest(runner.canonical(manifest))==expected
 for f in manifest['outputs']:
  h=hashlib.sha256();size=0
  for i in range((f['size_bytes']+32767)//32768):
   part=rpc('output-chunk',{'run_id':key['run_id'],'manifest_sha256':expected,'alias':f['alias'],'index':i})
   body=base64.b64decode(part['content_base64']);assert runner.digest(body)==part['sha256'];h.update(body);size+=len(body)
  assert size==f['size_bytes'] and h.hexdigest()==f['sha256']
  collected.append({'size_bytes':size,'sha256':h.hexdigest()})
print(json.dumps({'result':result['result'],'collected':collected}))
'''

def four_files():
    sizes=[2097153,2097153,2097153,2097149];body=b'a'*runner.LIMIT
    source="for i in range(4):\n data=open('/input/input_'+str(i),'rb').read();open('/output/result'+str(i)+'.csv','wb').write(data)"
    stage=fixture(body,source);d=stage['workbench']['descriptor']
    d['inputs']=[{'alias':'input_'+str(i),'file_id':f'33333333-3333-4333-8333-{i+1:012d}','name':f'input{i}.csv','size_bytes':size,'sha256':runner.digest(b'a'*size)} for i,size in enumerate(sizes)]
    d['outputs']=[{'alias':'output_'+str(i),'name':f'result{i}.csv','size_limit_bytes':size} for i,size in enumerate(sizes)]
    d['code']['input_aliases']=[f['alias'] for f in d['inputs']]
    d['code']['outputs']=[{'name':f['name'],'size_limit_bytes':f['size_limit_bytes']} for f in d['outputs']]
    stage['request']['descriptor_sha256']=runner.digest(runner.canonical(d))
    return body,stage

@unittest.skipUnless(IMAGE,'AX_CODE_IMAGE required')
class EightMiBTests(unittest.TestCase):
    def invoke(self,content,source,stage=None):
        stage=stage or fixture(content,source)
        args=['docker','run','--rm','-i','--network','none','--read-only','--pids-limit','64','--memory','384m','--cpus','1','--cap-drop','ALL','--security-opt','no-new-privileges']
        for cap in ('KILL','SYS_CHROOT','SETUID','SETGID','SETPCAP'):args+=['--cap-add',cap]
        for path,options in [('/sandbox/input','size=8404992,nr_inodes=16,uid=0,gid=0,mode=755'),('/sandbox/tmp','size=8m,nr_inodes=128,uid=65532,gid=65532,mode=700'),('/sandbox/output','size=8400896,nr_inodes=32,uid=65532,gid=65532,mode=700'),('/var/lib/ax-code','size=256k,nr_inodes=32,uid=0,gid=0,mode=700')]:args+=['--tmpfs',path+':nosuid,nodev,noexec,'+options]
        args+=['--entrypoint','python3',IMAGE,'-I','-B','-c',DRIVER]
        p=subprocess.run(args,input=json.dumps({'stage':stage,'content':base64.b64encode(content).decode()}).encode(),capture_output=True,timeout=90)
        self.assertEqual(p.returncode,0,p.stderr.decode()[-2000:]);return json.loads(p.stdout)
    def test_exact_8m_input_output_chunk_roundtrip(self):
        content=b'a'*(8*1024*1024)
        result=self.invoke(content,"data=open('/input/input_1','rb').read();open('/output/result.csv','wb').write(data);print(len(data))")
        self.assertEqual(result['result']['status'],'succeeded',result)
        self.assertEqual(result['result']['summary'],'8388608\n')
        self.assertEqual(result['collected'],[{'size_bytes':len(content),'sha256':runner.digest(content)}])
    def test_csv_aggregate_and_xlsx_library(self):
        source="""import csv
from openpyxl import Workbook,load_workbook
from openpyxl.xml import DEFUSEDXML
assert DEFUSEDXML
amount=sum(int(row['amount']) for row in csv.DictReader(open('/input/input_1')))
w=Workbook();w.active.append(['total',amount]);w.save('/tmp/result.xlsx')
assert load_workbook('/tmp/result.xlsx',read_only=True).active['B1'].value==amount
open('/output/result.csv','w').write('total\\n'+str(amount)+'\\n')
print(amount)
"""
        content=b'amount\n2\n3\n';stage=fixture(content,source);d=stage['workbench']['descriptor']
        d['definition_manifest']=[{'id':'44444444-4444-4444-8444-444444444444','sha256':'a'*64,'size_bytes':1,'kind':'skill'}]
        d['history']=[{'kind':'user_start','text':'Aggregate'}]+[{'kind':'code_result','text':'prior result'} for _ in range(16)]
        stage['request']['descriptor_sha256']=runner.digest(runner.canonical(d))
        result=self.invoke(content,source,stage)
        self.assertEqual(result['result']['status'],'succeeded',result)
        self.assertEqual(result['result']['summary'],'5\n')
    def test_four_unaligned_files_exact_logical_limit(self):
        content,stage=four_files()
        result=self.invoke(content,stage['workbench']['descriptor']['code']['source'],stage)
        self.assertEqual(result['result']['status'],'succeeded',result)
        self.assertEqual(sum(f['size_bytes'] for f in result['collected']),runner.LIMIT)
    def test_signal_exit_is_nonnegative_and_no_outputs(self):
        result=self.invoke(b'x','while True: pass')
        self.assertEqual(result['result']['status'],'failed',result)
        self.assertGreaterEqual(result['result']['exit_code'],128)
        self.assertEqual(result['collected'],[])

if __name__=='__main__':unittest.main()

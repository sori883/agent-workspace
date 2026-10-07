import base64
import hashlib
import json
import os
from pathlib import Path
import subprocess
import unittest


IMAGE = os.environ.get("AX_CODE_IMAGE")


@unittest.skipUnless(IMAGE, "AX_CODE_IMAGE is required for Linux isolation tests")
class CodeRuntimeTests(unittest.TestCase):
    def run_code(self, source, inputs=None, outputs=None, omit_cap=None, omit_mount=None, extra=None, driver=None, launcher_args=None):
        args = ["docker", "run", "--rm", "-i", "--network", "none", "--read-only", "--pids-limit", "16", "--memory", "384m", "--cpus", "1", "--cap-drop", "ALL", "--security-opt", "no-new-privileges"]
        for cap in ("KILL", "SYS_CHROOT", "SETUID", "SETGID", "SETPCAP"):
            if cap != omit_cap:
                args += ["--cap-add", cap]
        mounts = {"input": "size=64k,nr_inodes=16,uid=0,gid=0,mode=755", "tmp": "size=8m,nr_inodes=128,uid=65532,gid=65532,mode=700", "output": "size=1m,nr_inodes=32,uid=65532,gid=65532,mode=700"}
        for name, options in mounts.items():
            if name != omit_mount:
                args += ["--tmpfs", "/sandbox/" + name + ":nosuid,nodev,noexec," + options]
        args += ["--tmpfs", "/var/lib/ax-code:nosuid,nodev,noexec,size=64k,nr_inodes=32,mode=700"]
        args += extra or []
        args += ["--entrypoint", "python3", IMAGE, "-I", "-B"]
        args += ["-c", driver] if driver else ["/opt/ax-code/launcher.py"]
        args += launcher_args or []
        request = {"version": 1, "source": source, "inputs": {name: base64.b64encode(data).decode() for name, data in (inputs or {}).items()}, "outputs": outputs or ["result.txt"]}
        process = subprocess.run(args, input=json.dumps(request).encode(), stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=20)
        self.assertEqual(process.returncode, 0, process.stderr.decode()[:1000])
        try:
            return json.loads(process.stdout)
        except ValueError:
            self.fail("launcher did not return structured result: " + process.stderr.decode()[:1000])

    def artifact(self, result, name="result.txt"):
        self.assertEqual(result["status"], "succeeded", result)
        self.assertTrue(result["isolation_ready"])
        item = next(value for value in result["outputs"] if value["name"] == name)
        body = base64.b64decode(item["content_base64"])
        self.assertEqual(hashlib.sha256(body).hexdigest(), item["sha256"])
        self.assertEqual(len(body), item["size_bytes"])
        return body

    def test_csv_aggregation(self):
        result = self.run_code("import csv\nfrom decimal import Decimal\nrows=list(csv.DictReader(open('/input/sales.csv')))\namount=sum(Decimal(r['amount']) for r in rows)\nopen('/output/result.txt','w').write(str(amount))", {"sales.csv": b"department,amount\nA,12.25\nA,7.75\nB,4.50\n"})
        self.assertEqual(self.artifact(result), b"24.50")

    def test_xlsx_roundtrip_and_defusedxml(self):
        source = """from openpyxl import Workbook,load_workbook
from openpyxl.xml import DEFUSEDXML
from defusedxml.ElementTree import fromstring
assert DEFUSEDXML
try: fromstring('<!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>')
except Exception: pass
else: raise AssertionError('entity accepted')
w=Workbook(); s=w.active; s.append(['department','amount']); s.append(['A',20]); s.append(['B',4.5]); w.save('/tmp/input.xlsx')
r=load_workbook('/tmp/input.xlsx',read_only=True,data_only=True)
assert sum(row[1] for row in list(r.active.values)[1:])==24.5
w=Workbook(); w.active.append(['total',24.5]); w.save('/output/result.xlsx')
assert load_workbook('/output/result.xlsx',read_only=True).active['B1'].value==24.5
"""
        self.assertTrue(self.artifact(self.run_code(source, outputs=["result.xlsx"]), "result.xlsx").startswith(b"PK"))

    def test_network_fork_threads_exec_and_parent_control_denied(self):
        source = """import os,socket,ctypes,threading,subprocess,stat
def denied(call):
 try: call()
 except (OSError,RuntimeError): return
 raise AssertionError('not denied')
for family,kind in [(socket.AF_INET,socket.SOCK_STREAM),(socket.AF_INET,socket.SOCK_DGRAM),(socket.AF_INET6,socket.SOCK_STREAM),(socket.AF_UNIX,socket.SOCK_STREAM)]: denied(lambda: socket.socket(family,kind))
denied(os.fork)
denied(lambda: threading.Thread(target=lambda:None).start())
denied(lambda: os.execve('/usr/local/bin/python3',['python3','-c','pass'],{}))
denied(lambda: os.kill(os.getppid(),0))
c=ctypes.CDLL(None,use_errno=True)
assert c.ptrace(0,0,0,0)==-1
assert c.prctl(38,0,0,0,0)==-1
assert c.mount(b'none',b'/tmp',b'tmpfs',0,0)==-1
assert c.setuid(0)==-1
assert c.mknod(b'/tmp/fifo',stat.S_IFIFO|0o600,0)==-1
assert c.mknodat(-100,b'/tmp/fifo-at',stat.S_IFIFO|0o600,0)==-1
open('/output/result.txt','w').write('blocked')
"""
        self.assertEqual(self.artifact(self.run_code(source)), b"blocked")

    def test_chroot_environment_readonly_inputs_and_fd_boundary(self):
        source = """import os
assert os.getuid()==65532 and os.getgid()==65532
assert 'SENTINEL_PARENT_ONLY' not in os.environ
for p in ['/proc/self/environ','/proc/1/root','/opt/ax-code/launcher.py','/var/lib/ax-code/started','/../../var/lib/ax-code/started']:
 try: open(p).read()
 except OSError: pass
 else: raise AssertionError('outside read')
for p in ['/input/source.txt','/trusted/security.py','/usr/local/lib/python3.12/os.py','/elsewhere']:
 try: open(p,'w').write('altered')
 except OSError: pass
 else: raise AssertionError('readonly write')
for n in range(3,64):
 try: os.fstat(n)
 except OSError: pass
 else: raise AssertionError('inherited fd')
open('/output/result.txt','w').write(open('/input/source.txt').read())
"""
        self.assertEqual(self.artifact(self.run_code(source, {"source.txt": b"intact"}, extra=["-e", "SENTINEL_PARENT_ONLY=outside"])), b"intact")

    def test_setup_failure_never_runs_code(self):
        for cap in ("KILL", "SYS_CHROOT", "SETUID", "SETGID", "SETPCAP"):
            with self.subTest(cap=cap):
                result = self.run_code("open('/output/result.txt','w').write('ran')", omit_cap=cap)
                self.assertEqual(result["status"], "setup_failed", result)
                self.assertFalse(result["isolation_ready"])
                self.assertEqual(result["outputs"], [])
        result = self.run_code("open('/output/result.txt','w').write('ran')", omit_mount="tmp")
        self.assertEqual(result["status"], "setup_failed", result)

    def test_seccomp_setup_failure_never_runs_code(self):
        profile = Path(__file__).resolve().parents[1] / "code_runtime/deny-seccomp-test.json"
        result = self.run_code("open('/output/result.txt','w').write('ran')", extra=["--security-opt", "seccomp=" + str(profile)])
        self.assertEqual(result["status"], "setup_failed", result)
        self.assertEqual(result.get("setup_detail"), "E:seccomp:0", result)
        self.assertFalse(result["isolation_ready"])

    def test_total_output_filesystem_quota(self):
        source = """import os,errno
try:
 for i in range(24):
  with open('/output/f'+str(i),'wb') as f: f.write(b'x'*65536)
except OSError as e: assert e.errno==errno.ENOSPC
else: raise AssertionError('total bytes unlimited')
for name in os.listdir('/output'): os.unlink('/output/'+name)
open('/output/result.txt','w').write('quota')
"""
        self.assertEqual(self.artifact(self.run_code(source)), b"quota")

    def test_supervisor_death_kills_running_child(self):
        driver = """import os,subprocess,sys,time,json,signal
p=subprocess.Popen(['python3','-I','-B','/opt/ax-code/launcher.py'],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
p.stdin.write(sys.stdin.buffer.read());p.stdin.close()
deadline=time.monotonic()+3
ready=False
while time.monotonic()<deadline:
 os.seteuid(65532)
 try: ready=os.path.exists('/sandbox/output/result.txt')
 finally: os.seteuid(0)
 if ready: break
 time.sleep(.01)
assert ready
child=int(open('/proc/'+str(p.pid)+'/task/'+str(p.pid)+'/children').read().strip())
os.kill(p.pid,signal.SIGKILL);p.wait(timeout=1)
deadline=time.monotonic()+1
while time.monotonic()<deadline:
 waited,status=os.waitpid(child,os.WNOHANG)
 if waited: break
 time.sleep(.01)
else: raise AssertionError('orphan still running')
assert os.WIFSIGNALED(status) and os.WTERMSIG(status)==signal.SIGKILL
print(json.dumps({'child_stopped':True}))
"""
        result = self.run_code("import time\nopen('/output/result.txt','w').write('ready')\nwhile True: time.sleep(1)", driver=driver)
        self.assertEqual(result, {"child_stopped": True})

    def test_wall_cpu_memory_and_log_limits(self):
        cases = [("import time\ntime.sleep(30)", "wall_limit"), ("while True: pass", "code_failed"), ("a=bytearray(1024**3)", "code_failed"), ("while True: print('x'*4096,flush=True)", "log_limit")]
        for source, code in cases:
            with self.subTest(code=code, source=source[:20]):
                result = self.run_code(source)
                self.assertTrue(result["isolation_ready"], result)
                self.assertEqual(result["status"], "failed", result)
                self.assertEqual(result["code"], code, result)
                self.assertLessEqual(len(base64.b64decode(result["untrusted_log_base64"])), 8192)

    def test_inode_fsize_and_descriptor_limits(self):
        source = """import os,errno
try:
 for i in range(200): open('/tmp/f'+str(i),'wb').close()
except OSError as e: assert e.errno==errno.ENOSPC
else: raise AssertionError('inode unlimited')
for name in os.listdir('/tmp'): os.unlink('/tmp/'+name)
try:
 with open('/tmp/big','wb') as f: f.write(b'x'*65537)
except OSError as e: assert e.errno==errno.EFBIG
else: raise AssertionError('size unlimited')
fds=[]
try:
 for i in range(100): fds.append(os.open('/input/main.py',os.O_RDONLY))
except OSError as e: assert e.errno==errno.EMFILE
else: raise AssertionError('fd unlimited')
for fd in fds: os.close(fd)
open('/output/result.txt','w').write('limited')
"""
        self.assertEqual(self.artifact(self.run_code(source)), b"limited")

    def test_symlink_hardlink_and_unexpected_output_rejected(self):
        cases = ["import os\nos.symlink('/input/main.py','/output/result.txt')", "import os\nopen('/output/other','w').write('a');os.link('/output/other','/output/result.txt')", "import os\nos.mkdir('/output/result.txt')"]
        for source in cases:
            with self.subTest(source=source):
                result = self.run_code(source)
                self.assertEqual(result["status"], "failed", result)
                self.assertEqual(result["outputs"], [])


if __name__ == "__main__":
    unittest.main()
